"""Regression checks for the verified-result delivery boundary."""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXECUTOR_PATH = (
    ROOT / "assets" / "portable-scenario-executor" / "scripts" / "execute_scenario.py"
)
SPEC = importlib.util.spec_from_file_location("portable_executor_delivery", EXECUTOR_PATH)
assert SPEC and SPEC.loader
EXECUTOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(EXECUTOR)


def verified_payload() -> dict:
    return {
        "status": "completed_deterministically",
        "request": "export reviewed duplicate charges",
        "selected_rule": {"source_id": "rules", "row": {"rule_id": "R-1"}},
        "recipe_execution": {
            "status": "verified_recipe",
            "verified": True,
            "recipe_id": "recipe-duplicate",
            "catalog_fingerprint": "a" * 64,
            "verification_path": "references/recipe-verification.json",
        },
        "deterministic_result": {
            "recipe_id": "recipe-duplicate",
            "columns": ["case_id", "amount", "decision"],
            "rows": [
                {"case_id": "C-1", "amount": 30, "decision": "duplicate"},
                {"case_id": "C-2", "amount": 20, "decision": "duplicate"},
            ],
            "summary": {"matched_row_count": 2, "matched_group_count": 2},
            "coverage": {
                "complete_for_all_matching_runtime_rows": True,
                "returned_row_count": 2,
                "total_matched_row_count": 2,
                "truncated": False,
            },
        },
        "result_handle": {"kind": "deterministic_result_handle", "result_id": "result-test"},
        "result_contract": {
            "required_fields": ["business_conclusion", "coverage"],
            "declared_output_contract": {},
            "design_time_templates": [
                {
                    "template_id": "csv-exact",
                    "name": "duplicate_rows",
                    "format": "csv",
                    "columns": ["case_id", "amount", "decision"],
                },
                {
                    "template_id": "csv-incomplete",
                    "name": "requires_agent_conclusion",
                    "format": "csv",
                    "columns": ["case_id", "business_conclusion"],
                },
                {
                    "template_id": "xlsx-exact",
                    "name": "duplicate:/?*[]\\rows",
                    "format": "xlsx",
                    "columns": ["case_id", "amount", "decision"],
                },
            ],
        },
    }


def structured_template(
    template_id: str, *, format: str = "xlsx", worksheet_name: str = "Historical Result",
    columns: list[str] | None = None, materialization_status: str = "materializable_schema",
    reason: str = "single_tabular_schema",
) -> dict:
    columns = columns or ["case_id", "amount", "decision"]
    return {
        "template_id": template_id,
        "name": "historical_result",
        "format": format,
        "historical_data_policy": "structure_metadata_only_no_historical_rows",
        "tables": [{
            "table_id": "table-1",
            "table_name": worksheet_name,
            "worksheet_name": worksheet_name if format == "xlsx" else None,
            "header_row_index": 0,
            "columns": [
                {
                    "ordinal": index,
                    "name": column,
                    "semantic_kind": "other",
                    "source_type": None,
                }
                for index, column in enumerate(columns)
            ],
        }],
        "materialization": {
            "status": materialization_status,
            "reason": reason,
            "message": "portable test structure",
            "selected_table_id": "table-1",
            "worksheet_name": worksheet_name if format == "xlsx" else None,
            "header_row_index": 0,
            "column_order": columns,
            "layout_fidelity": "schema_only",
        },
        "structure_fingerprint": "f" * 64,
    }


class PortableExecutorDeliveryTests(unittest.TestCase):
    def setUp(self) -> None:
        self._previous_artifact_root = os.environ.get("BUSINESS_ARTIFACT_ROOT")
        self._artifact_root = tempfile.TemporaryDirectory()
        os.environ["BUSINESS_ARTIFACT_ROOT"] = self._artifact_root.name

    def tearDown(self) -> None:
        if self._previous_artifact_root is None:
            os.environ.pop("BUSINESS_ARTIFACT_ROOT", None)
        else:
            os.environ["BUSINESS_ARTIFACT_ROOT"] = self._previous_artifact_root
        self._artifact_root.cleanup()

    def test_evidence_only_payload_never_creates_a_final_result_file(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            target = Path(self._artifact_root.name) / "not-a-final-result.json"
            payload = verified_payload()
            payload["status"] = "ready_for_agent_judgment"
            payload["recipe_execution"]["verified"] = False
            payload["recipe_execution"]["status"] = "unverified_recipe_evidence_only"

            delivery = EXECUTOR.materialize_verified_delivery(payload, target)

            self.assertEqual("blocked_delivery_requires_verified_result", delivery["status"])
            self.assertFalse(target.exists())

    def test_flow_contract_preserves_structured_template_for_delivery(self) -> None:
        template = structured_template("historical-xlsx", worksheet_name="Historical Outcome")
        template["output_columns"] = ["case_id", "amount", "decision"]
        contract = EXECUTOR.result_contract({"design_time_output_templates": [template]})

        resolved = contract["design_time_templates"][0]
        self.assertEqual("historical-xlsx", resolved["template_id"])
        self.assertEqual("Historical Outcome", resolved["tables"][0]["worksheet_name"])
        self.assertEqual("materializable_schema", resolved["materialization"]["status"])
        self.assertEqual("structure_metadata_only_no_historical_rows", resolved["historical_data_policy"])

    def test_saved_verified_result_can_materialize_json_without_reexecution(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            artifact_root = Path(self._artifact_root.name)
            evidence = artifact_root / "scenario-evidence.json"
            target = artifact_root / "result.json"
            evidence.write_text(json.dumps(verified_payload(), ensure_ascii=False), encoding="utf-8")

            delivery = EXECUTOR.run([
                "deliver", "--result", "scenario-evidence.json", "--output", "result.json", "--format", "json",
            ])

            self.assertEqual("delivered_deterministically", delivery["status"])
            self.assertTrue(target.is_file())
            exported = json.loads(target.read_text(encoding="utf-8"))
            self.assertEqual("verified_deterministic_result_delivery", exported["kind"])
            self.assertTrue(exported["verification"]["verified"])
            self.assertEqual(["case_id", "amount", "decision"], exported["deterministic_result"]["columns"])
            self.assertEqual("scenario-evidence.json", exported["evidence_package"]["relative_path"])
            self.assertFalse(exported["delivery_boundary"]["agent_semantic_judgment_used"])

    def test_delivery_never_overwrites_the_evidence_package(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            evidence = Path(self._artifact_root.name) / "scenario-evidence.json"
            original = json.dumps(verified_payload(), ensure_ascii=False)
            evidence.write_text(original, encoding="utf-8")

            delivery = EXECUTOR.run([
                "deliver", "--result", "scenario-evidence.json", "--output", "scenario-evidence.json",
            ])

            self.assertEqual("blocked_delivery_output_conflicts_with_evidence", delivery["status"])
            self.assertEqual(original, evidence.read_text(encoding="utf-8"))

    def test_csv_requires_an_exact_template_mapping_and_writes_a_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            evidence = root / "scenario-evidence.json"
            evidence.write_text(json.dumps(verified_payload(), ensure_ascii=False), encoding="utf-8")
            csv_target = root / "duplicate-rows.csv"

            delivery = EXECUTOR.run([
                "deliver", "--result", "scenario-evidence.json", "--output", "duplicate-rows.csv",
                "--template-id", "csv-exact",
            ])

            self.assertEqual("delivered_deterministically", delivery["status"])
            with csv_target.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(["case_id", "amount", "decision"], list(rows[0]))
            self.assertEqual("30", rows[0]["amount"])
            manifest_path = root / "duplicate-rows.delivery.json"
            self.assertTrue(manifest_path.is_file())
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual("materialized_exact_columns", manifest["template_materialization"]["status"])

            blocked_target = root / "must-not-exist.csv"
            blocked = EXECUTOR.run([
                "deliver", "--result", "scenario-evidence.json", "--output", "must-not-exist.csv",
                "--template-id", "csv-incomplete",
            ])
            self.assertEqual("blocked_delivery_template_not_materializable", blocked["status"])
            self.assertFalse(blocked_target.exists())

    def test_csv_formula_like_text_is_exported_as_literal_data(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            evidence = root / "scenario-evidence.json"
            payload = verified_payload()
            payload["deterministic_result"]["rows"][0]["decision"] = "\t=not-a-formula"
            evidence.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            target = root / "duplicate-rows.csv"

            delivery = EXECUTOR.run([
                "deliver", "--result", "scenario-evidence.json", "--output", "duplicate-rows.csv",
                "--template-id", "csv-exact",
            ])

            self.assertEqual("delivered_deterministically", delivery["status"])
            with target.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual("'\t=not-a-formula", rows[0]["decision"])

    def test_xlsx_requires_the_same_verified_template_mapping_and_is_data_only(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            evidence = root / "scenario-evidence.json"
            payload = verified_payload()
            payload["deterministic_result"]["rows"][0]["decision"] = "=not-a-formula"
            evidence.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            target = root / "duplicate-rows.xlsx"

            delivery = EXECUTOR.run([
                "deliver", "--result", "scenario-evidence.json", "--output", "duplicate-rows.xlsx",
                "--format", "xlsx", "--template-id", "xlsx-exact",
            ])

            self.assertEqual("delivered_deterministically", delivery["status"])
            self.assertTrue(target.is_file())
            descriptor = delivery["result_file"]
            self.assertEqual("verified_deterministic_result_workbook", descriptor["kind"])
            self.assertEqual("xlsx", descriptor["format"])
            self.assertEqual(hashlib.sha256(target.read_bytes()).hexdigest(), descriptor["sha256"])
            self.assertEqual(["case_id", "amount", "decision"], descriptor["columns"])
            self.assertEqual(2, descriptor["row_count"])
            self.assertLessEqual(len(descriptor["sheet_name"]), 31)
            self.assertFalse(any(character in descriptor["sheet_name"] for character in r"\\/*?:[]"))

            workbook = load_workbook(target, data_only=False, read_only=True)
            try:
                worksheet = workbook[descriptor["sheet_name"]]
                rows = list(worksheet.iter_rows())
                self.assertEqual(["case_id", "amount", "decision"], [cell.value for cell in rows[0]])
                self.assertEqual("C-1", rows[1][0].value)
                self.assertEqual(30, rows[1][1].value)
                self.assertEqual("'=not-a-formula", rows[1][2].value)
                self.assertEqual("s", rows[1][2].data_type)
            finally:
                workbook.close()
            self.assertTrue((root / "duplicate-rows.delivery.json").is_file())

    def test_single_materializable_historical_template_is_selected_automatically(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            evidence = root / "scenario-evidence.json"
            payload = verified_payload()
            payload["result_contract"]["design_time_templates"] = [
                structured_template("historical-xlsx", worksheet_name="Historical Outcome")
            ]
            evidence.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            target = root / "historical.xlsx"

            delivery = EXECUTOR.run([
                "deliver", "--result", "scenario-evidence.json", "--output", "historical.xlsx", "--format", "xlsx",
            ])

            self.assertEqual("delivered_deterministically", delivery["status"])
            materialization = delivery["template_materialization"]
            self.assertEqual("historical-xlsx", materialization["template_id"])
            self.assertEqual("auto_single_compatible_template", materialization["selection"])
            self.assertEqual("materialized_historical_structure", materialization["status"])
            self.assertEqual("Historical Outcome", delivery["result_file"]["sheet_name"])
            workbook = load_workbook(target, read_only=True, data_only=True)
            try:
                worksheet = workbook["Historical Outcome"]
                self.assertEqual(
                    ["case_id", "amount", "decision"],
                    [cell.value for cell in next(worksheet.iter_rows())],
                )
            finally:
                workbook.close()

    def test_multiple_compatible_historical_templates_require_explicit_selection(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            payload = verified_payload()
            payload["result_contract"]["design_time_templates"] = [
                structured_template("first", worksheet_name="First"),
                structured_template("second", worksheet_name="Second"),
            ]
            target = root / "must-not-exist.xlsx"

            delivery = EXECUTOR.materialize_verified_delivery(payload, target, "xlsx")

            self.assertEqual("blocked_delivery_template_selection_required", delivery["status"])
            self.assertEqual(["first", "second"], delivery["available_template_ids"])
            self.assertFalse(target.exists())

    def test_nonmaterializable_historical_template_blocks_without_lookalike_file(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            payload = verified_payload()
            payload["result_contract"]["design_time_templates"] = [
                structured_template(
                    "multi-sheet", materialization_status="not_materializable",
                    reason="multiple_or_missing_tables",
                )
            ]
            target = root / "must-not-exist.xlsx"

            delivery = EXECUTOR.materialize_verified_delivery(payload, target, "xlsx")

            self.assertEqual("blocked_delivery_template_not_materializable", delivery["status"])
            self.assertFalse(target.exists())

    def test_tampered_historical_worksheet_name_is_blocked_not_sanitized(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            payload = verified_payload()
            payload["result_contract"]["design_time_templates"] = [
                structured_template("bad-sheet", worksheet_name="Bad/Worksheet")
            ]
            target = root / "must-not-exist.xlsx"

            delivery = EXECUTOR.materialize_verified_delivery(payload, target, "xlsx", "bad-sheet")

            self.assertEqual("blocked_delivery_template_not_materializable", delivery["status"])
            self.assertEqual("unsupported_worksheet_name", delivery["reason"])
            self.assertFalse(target.exists())

    def test_explicit_structured_template_refuses_missing_recipe_columns(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(self._artifact_root.name)
            payload = verified_payload()
            payload["result_contract"]["design_time_templates"] = [
                structured_template("extra-column", columns=["case_id", "amount", "decision", "rule_basis"])
            ]
            target = root / "must-not-exist.xlsx"

            delivery = EXECUTOR.materialize_verified_delivery(payload, target, "xlsx", "extra-column")

            self.assertEqual("blocked_delivery_template_not_materializable", delivery["status"])
            self.assertEqual(["rule_basis"], delivery["missing_columns"])
            self.assertFalse(target.exists())

    def test_absolute_delivery_names_are_blocked_by_host_artifact_boundary(self) -> None:
        evidence = Path(self._artifact_root.name) / "scenario-evidence.json"
        evidence.write_text(json.dumps(verified_payload(), ensure_ascii=False), encoding="utf-8")

        delivery = EXECUTOR.run([
            "deliver", "--result", "scenario-evidence.json", "--output", "E:/outputs/result.xlsx",
        ])

        self.assertEqual("blocked_host_artifact_sink", delivery["status"])
        self.assertFalse((Path("E:/outputs") / "result.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
