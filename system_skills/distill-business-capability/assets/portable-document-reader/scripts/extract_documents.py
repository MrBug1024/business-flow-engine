#!/usr/bin/env python3
"""Build and query provenance-preserving indexes for portable business documents."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import re
import sqlite3
import sys
import tempfile
from pathlib import Path, PurePosixPath
from typing import Any, Iterable, Iterator, Sequence


SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".html", ".htm", ".xml", ".json", ".yaml", ".yml",
    ".docx", ".pptx", ".pdf",
}
DEFAULT_MAX_CHARS = 200_000
HARD_MAX_CHARS = 2_000_000
DEFAULT_CHUNK_CHARS = 4_000
MAX_CHUNK_CHARS = 20_000
MAX_RESULTS = 100
ARTIFACT_ROOT_ENV = "BUSINESS_ARTIFACT_ROOT"


class ExtractError(ValueError):
    pass


class ArtifactPathError(ExtractError):
    pass


def artifact_root() -> Path:
    configured = str(os.environ.get(ARTIFACT_ROOT_ENV, "")).strip()
    if not configured:
        raise ArtifactPathError(
            "Persistent artifact output is blocked: the host must set BUSINESS_ARTIFACT_ROOT."
        )
    root = Path(configured).expanduser()
    if not root.is_absolute():
        raise ArtifactPathError("BUSINESS_ARTIFACT_ROOT must be an absolute host-managed directory.")
    root = root.resolve()
    root.mkdir(parents=True, exist_ok=True)
    if not root.is_dir():
        raise ArtifactPathError("BUSINESS_ARTIFACT_ROOT is not a writable artifact directory.")
    return root


def normalize_artifact_name(value: str | Path, label: str = "artifact name") -> str:
    raw = str(value or "").strip()
    normalized = raw.replace("\\", "/")
    if not normalized:
        raise ArtifactPathError(f"{label} must be a non-empty relative artifact name.")
    if (
        "\x00" in normalized
        or normalized.startswith("/")
        or raw.startswith("\\")
        or re.match(r"^[A-Za-z]:", raw)
    ):
        raise ArtifactPathError(f"{label} must be a safe relative artifact name, not an absolute path.")
    parts = PurePosixPath(normalized).parts
    if not parts or any(part in {"", ".", ".."} or ":" in part for part in parts):
        raise ArtifactPathError(f"{label} must not contain traversal or drive-qualified segments.")
    return PurePosixPath(*parts).as_posix()


def _absolute_artifact_input(raw: str) -> bool:
    return (
        Path(raw).expanduser().is_absolute()
        or raw.replace("\\", "/").startswith("/")
        or raw.startswith("\\")
        or bool(re.match(r"^[A-Za-z]:", raw))
    )


def resolve_artifact_name(value: str | Path, label: str = "artifact name") -> tuple[Path, str]:
    root = artifact_root()
    raw = str(value or "").strip()
    if _absolute_artifact_input(raw):
        target = Path(raw).expanduser().resolve()
        try:
            relative_path = target.relative_to(root).as_posix()
        except ValueError as exc:
            raise ArtifactPathError(
                f"{label} must be a safe relative artifact name or an already-translated path inside BUSINESS_ARTIFACT_ROOT."
            ) from exc
    else:
        relative_path = normalize_artifact_name(raw, label)
        target = (root.joinpath(*PurePosixPath(relative_path).parts)).resolve()
    if target == root or root not in target.parents:
        raise ArtifactPathError(f"{label} escapes BUSINESS_ARTIFACT_ROOT.")
    return target, relative_path


def artifact_reference(path: Path, kind: str) -> dict[str, Any]:
    root = artifact_root()
    resolved = path.expanduser().resolve()
    try:
        relative_path = resolved.relative_to(root).as_posix()
    except ValueError as exc:
        raise ArtifactPathError("Artifact output must stay inside BUSINESS_ARTIFACT_ROOT.") from exc
    if not resolved.is_file():
        raise ArtifactPathError("Artifact file does not exist inside BUSINESS_ARTIFACT_ROOT.")
    sha256 = file_digest(resolved)
    return {
        "schema_version": 1,
        "kind": kind,
        "artifact_id": "artifact-" + hashlib.sha256(
            f"{kind}\0{relative_path}\0{sha256}".encode("utf-8")
        ).hexdigest()[:24],
        "relative_path": relative_path,
        "format": resolved.suffix.casefold().lstrip("."),
        "sha256": sha256,
        "size_bytes": resolved.stat().st_size,
    }


def artifact_relative_path(path: Path) -> tuple[Path, str]:
    root = artifact_root()
    resolved = path.expanduser().resolve()
    try:
        relative_path = resolved.relative_to(root).as_posix()
    except ValueError as exc:
        raise ArtifactPathError("Artifact output must stay inside BUSINESS_ARTIFACT_ROOT.") from exc
    if not relative_path or relative_path == ".":
        raise ArtifactPathError("Artifact output must name a file below BUSINESS_ARTIFACT_ROOT.")
    return resolved, relative_path


def blocked_artifact_output(exc: ArtifactPathError) -> dict[str, Any]:
    return {
        "status": "blocked_host_artifact_sink",
        "message": str(exc),
        "artifact_root_environment": ARTIFACT_ROOT_ENV,
    }


def validate_file(raw: str) -> Path:
    path = Path(raw).expanduser().resolve()
    if not path.is_file():
        raise ExtractError(f"File does not exist: {path}")
    if path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
        raise ExtractError(f"Unsupported document extension: {path.suffix or '<none>'}")
    return path


def normalize(text: str) -> str:
    return re.sub(r"[ \t]+", " ", str(text or "")).strip()


def strip_markup(text: str) -> str:
    without_hidden = re.sub(r"(?is)<(?:script|style)\b.*?</(?:script|style)>", " ", text)
    return normalize(html.unescape(re.sub(r"(?s)<[^>]+>", " ", without_hidden)))


def detect_text_encoding(path: Path) -> str:
    with path.open("rb") as handle:
        sample = handle.read(64 * 1024)
    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    try:
        sample.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        try:
            sample.decode("gb18030")
            return "gb18030"
        except UnicodeDecodeError:
            return "latin-1"


def split_segment(locator: str, text: str, chunk_chars: int) -> Iterator[tuple[str, str]]:
    cleaned = str(text or "").strip()
    if not cleaned:
        return
    for start in range(0, len(cleaned), chunk_chars):
        part = cleaned[start : start + chunk_chars]
        suffix = f";chars:{start + 1}-{start + len(part)}" if len(cleaned) > chunk_chars else ""
        yield locator + suffix, part


def iter_plain_segments(path: Path, chunk_chars: int) -> Iterator[tuple[str, str]]:
    extension = path.suffix.casefold()
    character_start = 1
    line_start = 1
    with path.open("r", encoding=detect_text_encoding(path), errors="replace") as handle:
        while raw := handle.read(chunk_chars):
            newlines = raw.count("\n")
            line_end = line_start + newlines
            value = strip_markup(raw) if extension in {".html", ".htm", ".xml"} else raw.strip()
            if value:
                yield (
                    f"lines:{line_start}-{line_end};chars:{character_start}-{character_start + len(raw) - 1}",
                    value,
                )
            character_start += len(raw)
            line_start = line_end


def iter_docx_segments(path: Path, chunk_chars: int) -> Iterator[tuple[str, str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractError("python-docx is required for DOCX") from exc
    document = Document(str(path))
    paragraph_index = 0
    table_index = 0
    blocks = document.iter_inner_content() if hasattr(document, "iter_inner_content") else [
        *document.paragraphs, *document.tables
    ]
    for block in blocks:
        if hasattr(block, "rows"):
            table_index += 1
            for row_index, row in enumerate(block.rows, 1):
                text = "\t".join(cell.text for cell in row.cells)
                yield from split_segment(f"table:{table_index};row:{row_index}", text, chunk_chars)
        else:
            paragraph_index += 1
            yield from split_segment(f"paragraph:{paragraph_index}", getattr(block, "text", ""), chunk_chars)


def iter_pptx_segments(path: Path, chunk_chars: int) -> Iterator[tuple[str, str]]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ExtractError("python-pptx is required for PPTX") from exc
    presentation = Presentation(str(path))
    for slide_number, slide in enumerate(presentation.slides, 1):
        parts = [str(getattr(shape, "text", "")).strip() for shape in slide.shapes]
        yield from split_segment(f"slide:{slide_number}", "\n".join(item for item in parts if item), chunk_chars)


def iter_pdf_segments(path: Path, chunk_chars: int) -> Iterator[tuple[str, str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractError("pypdf is required for PDF text extraction") from exc
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, 1):
        yield from split_segment(f"page:{page_number}", page.extract_text() or "", chunk_chars)


def iter_segments(path: Path, chunk_chars: int = DEFAULT_CHUNK_CHARS) -> Iterator[tuple[str, str]]:
    extension = path.suffix.casefold()
    if extension == ".docx":
        yield from iter_docx_segments(path, chunk_chars)
    elif extension == ".pptx":
        yield from iter_pptx_segments(path, chunk_chars)
    elif extension == ".pdf":
        yield from iter_pdf_segments(path, chunk_chars)
    else:
        yield from iter_plain_segments(path, chunk_chars)


def bounded_extract(path: Path, limit: int) -> tuple[str, dict[str, Any]]:
    parts: list[str] = []
    evidence: list[dict[str, Any]] = []
    total = 0
    truncated = False
    segment_count = 0
    first_locator = ""
    last_locator = ""
    for locator, text in iter_segments(path, min(DEFAULT_CHUNK_CHARS, limit)):
        segment_count += 1
        first_locator = first_locator or locator
        last_locator = locator
        separator_chars = 1 if parts else 0
        remaining = limit - total - separator_chars
        if remaining <= 0:
            truncated = True
            break
        addition = text[:remaining]
        parts.append(addition)
        evidence.append({
            "locator": locator,
            "text_digest": hashlib.sha256(addition.encode("utf-8")).hexdigest(),
        })
        total += separator_chars + len(addition)
        if len(addition) < len(text):
            truncated = True
            break
    text = "\n".join(parts)
    metadata = {
        "parser": "streaming-segment-extractor",
        "segment_count_read": segment_count,
        "first_locator": first_locator,
        "last_locator": last_locator,
        "evidence": evidence,
        "truncated": truncated,
    }
    if path.suffix.casefold() == ".pdf" and len(text.strip()) < 80:
        metadata["ocr_recommended"] = True
    return text, metadata


def file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def create_temporary_index(output: Path, source: str, source_digest: str, segments: Iterable[tuple[str, str]]) -> dict[str, Any]:
    """Create a process-private index for an executor's temporary evidence pass.

    This is deliberately distinct from a delivered index artifact. Its caller
    owns a ``TemporaryDirectory`` and must not return its path in host metadata.
    """
    output = output.expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_handle = tempfile.NamedTemporaryFile(
        prefix=f"{output.name}.", suffix=".tmp", dir=output.parent, delete=False
    )
    temporary = Path(temporary_handle.name)
    temporary_handle.close()
    connection = sqlite3.connect(temporary)
    count = 0
    characters = 0
    try:
        connection.executescript(
            """
            PRAGMA journal_mode=DELETE;
            CREATE TABLE metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE chunks(
                chunk_id INTEGER PRIMARY KEY,
                source TEXT NOT NULL,
                source_digest TEXT NOT NULL,
                locator TEXT NOT NULL,
                text TEXT NOT NULL,
                text_digest TEXT NOT NULL
            );
            CREATE INDEX idx_chunks_locator ON chunks(locator);
            """
        )
        for locator, text in segments:
            cleaned = str(text or "").strip()
            if not cleaned:
                continue
            count += 1
            characters += len(cleaned)
            connection.execute(
                "INSERT INTO chunks VALUES (?,?,?,?,?,?)",
                (count, source, source_digest, locator, cleaned, hashlib.sha256(cleaned.encode("utf-8")).hexdigest()),
            )
        for key, value in {
            "schema_version": "1",
            "source": source,
            "source_digest": source_digest,
            "chunk_count": str(count),
            "character_count": str(characters),
        }.items():
            connection.execute("INSERT INTO metadata VALUES (?,?)", (key, value))
        connection.commit()
    except Exception:
        connection.close()
        temporary.unlink(missing_ok=True)
        raise
    finally:
        try:
            connection.close()
        except sqlite3.Error:
            pass
    temporary.replace(output)
    return {
        "status": "success",
        "index": str(output),
        "source": source,
        "source_digest": source_digest,
        "chunk_count": count,
        "character_count": characters,
    }


def create_index(output: Path, source: str, source_digest: str, segments: Iterable[tuple[str, str]]) -> dict[str, Any]:
    """Create a persistent index only under the host-injected artifact root."""
    output, _relative_path = artifact_relative_path(output)
    result = create_temporary_index(output, source, source_digest, segments)
    artifact = artifact_reference(output, "document_evidence_index")
    result.pop("index", None)
    result.update({
        "artifact": artifact,
        "artifact_id": artifact["artifact_id"],
        "relative_path": artifact["relative_path"],
    })
    return result


def ocr_text_items(payload: Any, prefix: str = "ocr") -> Iterator[tuple[str, str, str]]:
    if isinstance(payload, list):
        for index, item in enumerate(payload, 1):
            yield from ocr_text_items(item, f"{prefix}:{index}")
    elif isinstance(payload, dict):
        source = str(payload.get("source") or payload.get("filename") or prefix)
        text = payload.get("text")
        if not isinstance(text, str):
            data = payload.get("data")
            text = data.get("text") if isinstance(data, dict) else None
        if isinstance(text, str) and text.strip():
            yield source, text, str(payload.get("source_digest", ""))
        else:
            for key, value in payload.items():
                if key not in {"text", "data"} and isinstance(value, (dict, list)):
                    yield from ocr_text_items(value, f"{prefix}:{key}")


def create_ocr_index(input_json: Path, output: Path, source_override: str, chunk_chars: int) -> dict[str, Any]:
    try:
        payload = json.loads(input_json.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ExtractError(f"Invalid OCR JSON: {exc}") from exc
    items = list(ocr_text_items(payload))
    if not items:
        raise ExtractError("OCR JSON contains no successful text")
    source_path = Path(source_override).expanduser().resolve() if source_override else None
    if source_path and source_path.is_file():
        source = str(source_path)
        source_digest = file_digest(source_path)
        digest_kind = "original_input_sha256"
    else:
        source = source_override or ", ".join(dict.fromkeys(item[0] for item in items))
        digests = {item[2] for item in items if item[2]}
        source_digest = next(iter(digests)) if len(digests) == 1 else file_digest(input_json)
        digest_kind = "original_input_sha256" if len(digests) == 1 else "ocr_json_sha256_fallback"

    def segments() -> Iterator[tuple[str, str]]:
        for item_index, (item_source, text, _digest) in enumerate(items, 1):
            pages = re.split(r"(?m)^\s*\[?(?:Page|页)\s*(\d+)\]?\s*$", text)
            if len(pages) > 1:
                if pages[0].strip():
                    yield from split_segment(f"ocr:{item_index};source:{item_source};preamble", pages[0], chunk_chars)
                for index in range(1, len(pages), 2):
                    page = pages[index]
                    body = pages[index + 1] if index + 1 < len(pages) else ""
                    yield from split_segment(f"ocr:{item_index};source:{item_source};page:{page}", body, chunk_chars)
            else:
                yield from split_segment(f"ocr:{item_index};source:{item_source}", text, chunk_chars)

    result = create_index(output, source, source_digest, segments())
    result["source_digest_kind"] = digest_kind
    return result


def open_index(raw: str) -> sqlite3.Connection:
    path = Path(raw).expanduser().resolve()
    if not path.is_file():
        raise ExtractError(f"Index does not exist: {path}")
    connection = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    return connection


def search_index(raw: str, terms: Sequence[str], limit: int) -> dict[str, Any]:
    normalized = list(dict.fromkeys(item.strip() for item in terms if item.strip()))
    if not normalized:
        raise ExtractError("At least one non-empty --term is required")
    connection = open_index(raw)
    try:
        predicates = " OR ".join("instr(lower(text), lower(?)) > 0" for _ in normalized)
        score = " + ".join("CASE WHEN instr(lower(text), lower(?)) > 0 THEN 1 ELSE 0 END" for _ in normalized)
        parameters = [*normalized, *normalized, limit]
        rows = connection.execute(
            f"SELECT chunk_id, source, source_digest, locator, text, text_digest, ({score}) AS score "
            f"FROM chunks WHERE {predicates} ORDER BY score DESC, chunk_id LIMIT ?",
            parameters,
        ).fetchall()
        return {
            "status": "success",
            "terms": normalized,
            "hits": [dict(row) for row in rows],
            "hit_count_returned": len(rows),
            "evidence_contract": "Every hit retains source_digest, locator, chunk_id, and text_digest.",
        }
    finally:
        connection.close()


def get_chunk(raw: str, chunk_id: int) -> dict[str, Any]:
    connection = open_index(raw)
    try:
        row = connection.execute(
            "SELECT chunk_id, source, source_digest, locator, text, text_digest FROM chunks WHERE chunk_id=?",
            (chunk_id,),
        ).fetchone()
        if row is None:
            raise ExtractError(f"Unknown chunk_id: {chunk_id}")
        return {"status": "success", "evidence": dict(row)}
    finally:
        connection.close()


def get_context(raw: str, chunk_id: int, before: int, after: int) -> dict[str, Any]:
    connection = open_index(raw)
    try:
        anchor = connection.execute("SELECT source FROM chunks WHERE chunk_id=?", (chunk_id,)).fetchone()
        if anchor is None:
            raise ExtractError(f"Unknown chunk_id: {chunk_id}")
        lower = max(1, chunk_id - before)
        upper = chunk_id + after
        rows = connection.execute(
            "SELECT chunk_id, source, source_digest, locator, text, text_digest FROM chunks "
            "WHERE source=? AND chunk_id BETWEEN ? AND ? ORDER BY chunk_id",
            (anchor[0], lower, upper),
        ).fetchall()
        return {
            "status": "success", "anchor_chunk_id": chunk_id,
            "evidence": [dict(row) for row in rows],
            "section_boundary_policy": (
                "These are bounded adjacent chunks. The Agent must use headings/locators to confirm a complete section; "
                "absence of a boundary is uncertainty, not permission to invent missing text."
            ),
        }
    finally:
        connection.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("inspect", "extract"):
        command = commands.add_parser(name)
        command.add_argument("--input", required=True)
        command.add_argument("--max-chars", type=int, default=DEFAULT_MAX_CHARS)
        command.add_argument("--format", choices=["text", "json"], default="json")
    index = commands.add_parser("index")
    index.add_argument("--input", required=True)
    index.add_argument("--output", required=True, help="Safe relative artifact name under BUSINESS_ARTIFACT_ROOT.")
    index.add_argument("--chunk-chars", type=int, default=DEFAULT_CHUNK_CHARS)
    index_ocr = commands.add_parser("index-ocr")
    index_ocr.add_argument("--input-json", required=True)
    index_ocr.add_argument("--output", required=True, help="Safe relative artifact name under BUSINESS_ARTIFACT_ROOT.")
    index_ocr.add_argument("--source", default="")
    index_ocr.add_argument("--chunk-chars", type=int, default=DEFAULT_CHUNK_CHARS)
    search = commands.add_parser("search")
    search.add_argument("--index", required=True)
    search.add_argument("--term", action="append", default=[])
    search.add_argument("--limit", type=int, default=20)
    get = commands.add_parser("get")
    get.add_argument("--index", required=True)
    get.add_argument("--chunk-id", type=int, required=True)
    context = commands.add_parser("context")
    context.add_argument("--index", required=True)
    context.add_argument("--chunk-id", type=int, required=True)
    context.add_argument("--before", type=int, default=2)
    context.add_argument("--after", type=int, default=2)
    return parser


def run(argv: Sequence[str] | None = None) -> tuple[dict[str, Any], str]:
    args = build_parser().parse_args(argv)
    if args.command == "search":
        try:
            index_path, _index_relative_path = resolve_artifact_name(args.index, "--index")
        except ArtifactPathError as exc:
            payload = blocked_artifact_output(exc)
        else:
            payload = search_index(str(index_path), args.term, max(1, min(args.limit, MAX_RESULTS)))
    elif args.command == "get":
        try:
            index_path, _index_relative_path = resolve_artifact_name(args.index, "--index")
        except ArtifactPathError as exc:
            payload = blocked_artifact_output(exc)
        else:
            payload = get_chunk(str(index_path), args.chunk_id)
    elif args.command == "context":
        try:
            index_path, _index_relative_path = resolve_artifact_name(args.index, "--index")
        except ArtifactPathError as exc:
            payload = blocked_artifact_output(exc)
        else:
            payload = get_context(
                str(index_path), args.chunk_id,
                max(0, min(args.before, 5)), max(0, min(args.after, 5)),
            )
    elif args.command == "index-ocr":
        input_json = Path(args.input_json).expanduser().resolve()
        chunk_chars = max(200, min(args.chunk_chars, MAX_CHUNK_CHARS))
        try:
            output_path, _output_relative_path = resolve_artifact_name(args.output, "--output")
        except ArtifactPathError as exc:
            payload = blocked_artifact_output(exc)
        else:
            payload = create_ocr_index(input_json, output_path, args.source, chunk_chars)
            artifact = artifact_reference(output_path, "document_evidence_index")
            payload.pop("index", None)
            payload.update({
                "artifact": artifact,
                "artifact_id": artifact["artifact_id"],
                "relative_path": artifact["relative_path"],
            })
    else:
        path = validate_file(args.input)
        limit = max(1, min(int(getattr(args, "max_chars", DEFAULT_MAX_CHARS)), HARD_MAX_CHARS))
        if args.command == "index":
            chunk_chars = max(200, min(args.chunk_chars, MAX_CHUNK_CHARS))
            try:
                output_path, _output_relative_path = resolve_artifact_name(args.output, "--output")
            except ArtifactPathError as exc:
                payload = blocked_artifact_output(exc)
            else:
                payload = create_index(output_path, str(path), file_digest(path), iter_segments(path, chunk_chars))
                artifact = artifact_reference(output_path, "document_evidence_index")
                payload.pop("index", None)
                payload.update({
                    "artifact": artifact,
                    "artifact_id": artifact["artifact_id"],
                    "relative_path": artifact["relative_path"],
                })
        else:
            text, metadata = bounded_extract(path, limit if args.command == "extract" else min(limit, 20_000))
            payload = {
                "status": "success",
                "source": str(path),
                "extension": path.suffix.casefold(),
                "size_bytes": path.stat().st_size,
                "character_count_returned": len(text),
                **metadata,
            }
            if args.command == "extract":
                payload["text"] = text
            if metadata.get("ocr_recommended"):
                payload["status"] = "ocr_recommended"
                payload["message"] = "The PDF text layer is sparse; invoke the scenario OCR Skill, then index its JSON output."
    rendered = payload.get("text", "") if getattr(args, "format", "json") == "text" and args.command == "extract" else json.dumps(payload, ensure_ascii=False, indent=2)
    return payload, rendered


def main(argv: Sequence[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        payload, rendered = run(argv)
        code = 0
    except (ExtractError, OSError, sqlite3.Error, ValueError) as exc:
        payload = {"status": "error", "message": str(exc)}
        rendered = json.dumps(payload, ensure_ascii=False, indent=2)
        code = 2
    print(rendered)
    return code


if __name__ == "__main__":
    raise SystemExit(main())
